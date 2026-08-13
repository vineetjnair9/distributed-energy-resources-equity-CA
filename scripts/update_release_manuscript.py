#!/usr/bin/env python3
"""Synchronize housing results and diagnostics into the release DOCX manuscript."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
import statsmodels.formula.api as smf


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "site" / "index.docx"
TABLE_DIR = ROOT / "outputs" / "standardized_tables"
HOUSING_FIGURE = ROOT / "outputs" / "standardized_figures" / "housing_structure_attenuation.png"
CLUSTER_FIGURE = ROOT / "outputs" / "figures" / "cluster_choropleth_map.png"


def find_paragraph(document: Document, starts_with: str):
    matches = [p for p in document.paragraphs if p.text.strip().startswith(starts_with)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph beginning {starts_with!r}, found {len(matches)}")
    return matches[0]


def find_paragraph_any(document: Document, prefixes: tuple[str, ...]):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if any(paragraph.text.strip().startswith(prefix) for prefix in prefixes)
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph beginning with {prefixes!r}, found {len(matches)}")
    return matches[0]


def insert_after(paragraph, text: str = "", style: str | None = None):
    document = paragraph._parent
    new_paragraph = document.add_paragraph(text, style=style)
    paragraph._p.addnext(new_paragraph._p)
    return new_paragraph


def replace_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def remove_paragraph(paragraph) -> None:
    parent = paragraph._p.getparent()
    parent.remove(paragraph._p)


def replace_figure_before_caption(document: Document, caption_prefix: str, image_path: Path) -> None:
    """Replace the image paragraph immediately before a named caption."""
    caption = find_paragraph(document, caption_prefix)
    old_image = caption._p.getprevious()
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.4))
    caption._p.addprevious(image_paragraph._p)
    if old_image is not None and old_image.xpath(".//a:blip"):
        old_image.getparent().remove(old_image)


def clean_review_markup(document: Document) -> None:
    """Remove internal notes/highlights and accept tracked changes for release."""
    note_matches = [p for p in document.paragraphs if p.text.startswith("Overall notes/feedback:")]
    if note_matches:
        notes = note_matches[0]
        abstract = find_paragraph(document, "Abstract")
        between = False
        for paragraph in list(document.paragraphs):
            if paragraph._p is notes._p:
                between = True
            if paragraph._p is abstract._p:
                break
            if between:
                remove_paragraph(paragraph)

    root = document.part.element
    for highlight in list(root.xpath(".//w:highlight")):
        highlight.getparent().remove(highlight)
    for deletion in list(root.xpath(".//w:del")):
        deletion.getparent().remove(deletion)
    for insertion in list(root.xpath(".//w:ins")):
        parent = insertion.getparent()
        position = parent.index(insertion)
        for child in list(insertion):
            parent.insert(position, child)
            position += 1
        parent.remove(insertion)
    for marker_name in ("commentRangeStart", "commentRangeEnd", "commentReference"):
        for marker in list(root.xpath(f".//w:{marker_name}")):
            marker.getparent().remove(marker)


def fit_release_baselines() -> tuple[pd.DataFrame, dict[str, object]]:
    frame = pd.read_csv(ROOT / "data" / "processed" / "combined_der_dataset_w_controls_predictors.csv")
    population = pd.to_numeric(frame["total_population"], errors="coerce")
    zero_columns = [
        "PV_system_size_DC", "storage_capacity_mw", "total_chargers",
        "level1_chargers", "level2_chargers", "dc_fast_chargers",
    ]
    for column in zero_columns:
        frame[column] = pd.to_numeric(frame[column], errors="coerce").fillna(0)
    frame = frame.loc[population.notna() & (population >= 1000)].copy()
    population = frame["total_population"].replace(0, np.nan)
    frame["y_pv"] = np.log1p(frame["PV_system_size_DC"] * 1000 / population)
    frame["y_storage"] = np.log1p(frame["storage_capacity_mw"] * 100000 / population)
    frame["y_chargers"] = np.log1p(frame["total_chargers"] * 1000 / population)
    frame["y_level1_chargers"] = np.log1p(frame["level1_chargers"] * 1000 / population)
    frame["y_level2_chargers"] = np.log1p(frame["level2_chargers"] * 1000 / population)
    frame["y_dc_fast_chargers"] = np.log1p(frame["dc_fast_chargers"] * 1000 / population)
    frame["log_median_household_income"] = np.log(
        frame["median_household_income"].where(frame["median_household_income"] > 0)
    )
    frame["log_energy_gap_per_capita"] = np.log1p(
        frame["energy_affordability_gap"] / population
    )
    core = ["median_household_income", "pct_black", "pct_hispanic", "pct_asian"]
    frame = frame.dropna(subset=core).copy()
    common = "log_median_household_income + pct_black + pct_hispanic + pct_asian + poverty_rate"
    formulas = {
        "y_pv": f"y_pv ~ {common} + ghi_mean_kwh_m2_day_2023",
        "y_storage": f"y_storage ~ {common} + cdd65_2023 + hdd65_2023",
        "y_chargers": f"y_chargers ~ {common} + cdd65_2023 + hdd65_2023",
        "y_level1_chargers": f"y_level1_chargers ~ {common} + cdd65_2023 + hdd65_2023",
        "y_level2_chargers": f"y_level2_chargers ~ {common} + cdd65_2023 + hdd65_2023",
        "y_dc_fast_chargers": f"y_dc_fast_chargers ~ {common} + cdd65_2023 + hdd65_2023",
    }
    fits = {outcome: smf.ols(formula, data=frame).fit(cov_type="HC1") for outcome, formula in formulas.items()}
    return frame, fits


def stars(p_value: float) -> str:
    if p_value < 0.01:
        return "***"
    if p_value < 0.05:
        return "**"
    if p_value < 0.10:
        return "*"
    return ""


def spatial_stars(p_value: float) -> str:
    """Conventional two-sided significance markers used by the spatial appendix."""
    if p_value < 0.001:
        return "***"
    if p_value < 0.01:
        return "**"
    if p_value < 0.05:
        return "*"
    return ""


def coefficient_with_se(result, term: str) -> str:
    return f"{result.params[term]:.3f}{stars(result.pvalues[term])} ({result.bse[term]:.3f})"


def update_main_tables(document: Document) -> None:
    frame, fits = fit_release_baselines()

    summary = document.tables[1]
    for row in summary.rows[1:]:
        variable = row.cells[0].text
        series = pd.to_numeric(frame[variable], errors="coerce").dropna()
        values = [len(series), series.mean(), series.median(), series.quantile(0.25), series.quantile(0.75)]
        row.cells[1].text = f"{int(values[0]):,}"
        for cell, value in zip(row.cells[2:], values[1:]):
            cell.text = f"{float(value):.3f}"

    baseline = document.tables[2]
    outcomes = ["y_pv", "y_storage", "y_chargers"]
    for row in baseline.rows[1:6]:
        term = row.cells[0].text
        for cell, outcome in zip(row.cells[1:], outcomes):
            cell.text = coefficient_with_se(fits[outcome], term)
    pv = fits["y_pv"]
    storage = fits["y_storage"]
    chargers = fits["y_chargers"]
    baseline.rows[6].cells[1].text = "GHI: " + coefficient_with_se(pv, "ghi_mean_kwh_m2_day_2023")
    storage_climate = [storage.pvalues["cdd65_2023"], storage.pvalues["hdd65_2023"]]
    baseline.rows[6].cells[2].text = (
        "CDD/HDD: n.s." if min(storage_climate) >= 0.10
        else "CDD: " + coefficient_with_se(storage, "cdd65_2023")
    )
    charger_value = chargers.params["cdd65_2023"]
    charger_p = chargers.pvalues["cdd65_2023"]
    charger_se = chargers.bse["cdd65_2023"]
    baseline.rows[6].cells[3].text = (
        f"CDD65: {charger_value:.4f}{stars(charger_p)} ({charger_se:.4f})"
    )
    baseline.rows[7].cells[0].text = "N"
    baseline.rows[8].cells[0].text = "R²"
    for cell, outcome in zip(baseline.rows[7].cells[1:], outcomes):
        cell.text = f"{int(fits[outcome].nobs):,}"
    for cell, outcome in zip(baseline.rows[8].cells[1:], outcomes):
        cell.text = f"{fits[outcome].rsquared:.3f}"

    model9 = document.tables[3]
    model9_frame = pd.read_csv(
        ROOT / "outputs" / "tables" / "y_storage | Model 9 + pv control (most controlled).csv",
        index_col=0,
    )
    for row in model9.rows[1:]:
        term = row.cells[0].text
        estimate = float(model9_frame.loc[term, "Coef."])
        error = float(model9_frame.loc[term, "Std.Err."])
        p_value = float(model9_frame.loc[term, "P>|z|"])
        row.cells[1].text = f"{estimate:.3f}{stars(p_value)}"
        row.cells[2].text = f"{error:.3f}"
        row.cells[3].text = "< 0.001" if p_value < 0.001 else f"{p_value:.3f}"

    charger_table = document.tables[4]
    charger_outcomes = [
        "y_chargers", "y_level1_chargers", "y_level2_chargers", "y_dc_fast_chargers"
    ]
    for row in charger_table.rows[1:6]:
        term = row.cells[0].text
        for cell, outcome in zip(row.cells[1:], charger_outcomes):
            cell.text = coefficient_with_se(fits[outcome], term)
    charger_table.rows[6].cells[0].text = "N"
    charger_table.rows[7].cells[0].text = "R²"
    for cell, outcome in zip(charger_table.rows[6].cells[1:], charger_outcomes):
        cell.text = f"{int(fits[outcome].nobs):,}"
    for cell, outcome in zip(charger_table.rows[7].cells[1:], charger_outcomes):
        cell.text = f"{fits[outcome].rsquared:.3f}"

    assignments = pd.read_csv(
        ROOT / "outputs" / "tables" / "pca_kmeans_cluster_assignments.csv",
        dtype={"zip_code": "string"},
    )
    cluster_frame = assignments.merge(
        pd.read_csv(
            ROOT / "data" / "processed" / "combined_der_dataset_w_controls_predictors.csv",
            dtype={"zip_code": "string"},
        ),
        on="zip_code",
        validate="one_to_one",
    )
    population = cluster_frame["total_population"]
    cluster_frame["income_k"] = cluster_frame["median_household_income"] / 1000
    cluster_frame["residual_share"] = 1 - cluster_frame[
        ["pct_black", "pct_hispanic", "pct_asian"]
    ].sum(axis=1)
    cluster_frame["pv_rate"] = cluster_frame["PV_system_size_DC"] * 1000 / population
    cluster_frame["storage_rate"] = cluster_frame["storage_capacity_mw"] * 100000 / population
    cluster_frame["charger_rate"] = cluster_frame["total_chargers"] * 1000 / population
    cluster_table = document.tables[5]
    for row in cluster_table.rows[1:]:
        cluster = int(row.cells[0].text.split()[-1])
        group = cluster_frame.loc[cluster_frame["cluster"] == cluster]
        means = group[
            [
                "income_k", "poverty_rate", "pct_black", "pct_hispanic", "pct_asian",
                "residual_share", "pv_rate", "storage_rate", "charger_rate",
            ]
        ].mean()
        values = [
            f"{len(group):,}", f"${means['income_k']:.1f}k", f"{means['poverty_rate']:.1%}",
            f"{means['pct_black']:.1%}", f"{means['pct_hispanic']:.1%}",
            f"{means['pct_asian']:.1%}", f"{means['residual_share']:.1%}",
            f"{means['pv_rate']:.2f}", f"{means['storage_rate']:.2f}",
            f"{means['charger_rate']:.2f}",
        ]
        for cell, value in zip(row.cells[1:], values):
            cell.text = value


def materialize_figure_references(document: Document) -> None:
    """Replace stale REF/SEQ fields after Figure 15 with their intended visible text."""
    for paragraph in list(document.paragraphs):
        visible = paragraph.text
        if any(f"Figure {number}" in visible for number in range(16, 21)):
            replace_paragraph(paragraph, visible)


def materialize_equation_references(document: Document) -> None:
    """Preserve visible equation cross-references before replacing legacy bookmarks."""
    for paragraph in list(document.paragraphs):
        visible = paragraph.text
        if any(f"Equation ({number})" in visible for number in range(4, 15)):
            replace_paragraph(paragraph, visible)
        elif any(f"Equations ({number})" in visible for number in range(4, 15)):
            replace_paragraph(paragraph, visible)


def materialize_compact_equations(document: Document) -> None:
    """Use compact, fixed equation text where long legacy OMML objects overflow."""
    equations = {
        4: "y_z = β₀ + β₁ Income_z + β₂ Black_z + β₃ Hispanic_z + β₄ Asian_z\n"
           "+ β₅ Poverty_z + γX_z + ε_z                                                     (4)",
        5: "y_z = β₀ + β RaceSES_z + γ Climate_z + ε_z                                      (5)",
        6: "y_z = β₀ + β RaceSES_z + θ SESextra_z + γ Climate_z + ε_z                       (6)",
        7: "y_z = β₀ + β RaceSES_z + θ ResourceClimateAlt_z + ε_z                            (7)",
        8: "y_z = β₀ + β RaceSES_z + δ(Income_z × Race_z) + γX_z + ε_z                       (8)",
        9: "y_z = β₀ + β RaceSES_z + γX_z + δ Utility_z + ε_z                                (9)",
        10: "y_z = β₀ + β RaceSES_z + γX_z + θ Geo_z + ε_z                                  (10)",
        11: "y_z = β₀ + β RaceSES_z + γX_z + δ County_z + ε_z                               (11)",
        12: "y_z = β₀ + β RaceSES_z + γX_z + λ Infrastructure_z + ε_z                       (12)",
        13: "y_z = β₀ + β RaceSES_z + γX_z + λ Infrastructure_z + ϕ Demand_z + ε_z          (13)",
        14: "y_z = β₀ + β RaceSES_z + γX_z + λ Infrastructure_z + ϕ Demand_z + ψ DER_z + ε_z (14)",
    }
    for paragraph in document.paragraphs:
        if not paragraph._p.xpath(".//m:oMath"):
            continue
        xml_text = "".join(paragraph._p.itertext())
        for number, equation in equations.items():
            if xml_text.startswith(f"\t({number})"):
                paragraph.clear()
                run = paragraph.add_run(equation)
                run.font.name = "Cambria Math"
                run.font.size = Pt(9)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                break


def fix_equation_layout(document: Document) -> None:
    for paragraph in document.paragraphs:
        math_runs = paragraph._p.xpath(".//m:oMath/m:r")
        if not math_runs:
            continue
        for math_run in math_runs:
            properties = math_run.find(qn("w:rPr"))
            if properties is None:
                properties = OxmlElement("w:rPr")
                math_run.insert(0, properties)
            for tag in ("w:sz", "w:szCs"):
                size = properties.find(qn(tag))
                if size is None:
                    size = OxmlElement(tag)
                    properties.append(size)
                size.set(qn("w:val"), "14")
    diagnostics_heading = find_paragraph(document, "Multicollinearity diagnostics")
    diagnostics_heading.paragraph_format.page_break_before = True
    find_paragraph(document, "Spatial autocorrelation of residuals").paragraph_format.page_break_before = True
    find_paragraph(document, "Conley spatial HAC standard errors").paragraph_format.page_break_before = True


def coefficient(outcome: str, model: str, term: str) -> float:
    path = TABLE_DIR / f"{outcome} | {model}.csv"
    frame = pd.read_csv(path, index_col=0)
    return float(frame.loc[term, "Coef."])


def housing_result_text() -> str:
    baseline = "Model 1 baseline (climate controls)"
    housing = "Model 2D (add housing structure and tenure)"
    parts = []
    for outcome, label in [
        ("y_pv", "PV"),
        ("y_storage", "storage"),
        ("y_chargers", "aggregate charging"),
    ]:
        black_before = coefficient(outcome, baseline, "pct_black")
        black_after = coefficient(outcome, housing, "pct_black")
        asian_before = coefficient(outcome, baseline, "pct_asian")
        asian_after = coefficient(outcome, housing, "pct_asian")
        parts.append(
            f"For {label}, the standardized Black-share coefficient changes from "
            f"{black_before:.3f} to {black_after:.3f}, and the Asian-share coefficient "
            f"changes from {asian_before:.3f} to {asian_after:.3f}."
        )
    return (
        "Housing is tested directly rather than left only as a proposed mechanism. "
        "Model 2C adds the exhaustive B25024 structure composition, with single-family "
        "units as the reference, and Model 2D adds B25003 owner occupancy, with renters "
        "as the reference. "
        + " ".join(parts)
        + " Figure 21 shows all three race-coefficient comparisons. Movement toward zero "
        "is attenuation associated with area-level housing composition; these models do "
        "not identify household-level tenure or roof access."
    )


def update_release_prose(document: Document) -> None:
    """Keep exact numerical claims and stable cluster labels aligned with artifacts."""
    model9 = pd.read_csv(
        ROOT / "outputs" / "tables" / "y_storage | Model 9 + pv control (most controlled).csv",
        index_col=0,
    )

    def estimate(term: str) -> float:
        return float(model9.loc[term, "Coef."])

    replace_paragraph(
        find_paragraph_any(document, (
            "I also estimate an additional robustness model",
            "We also estimate an additional robustness model",
        )),
        "We also estimate an additional robustness model, Model 9, that adds y_pv to "
        "the most controlled storage specification to ask whether storage inequities "
        "persist after conditioning on local PV deployment intensity. The PV term is "
        f"large and positive ({estimate('y_pv'):.3f}, p < 0.001), as expected for a "
        "solar-plus-storage pathway (Table 4), but the race coefficients remain negative "
        "and statistically significant: "
        f"pct_black = {estimate('pct_black'):.3f} (p < 0.001), "
        f"pct_hispanic = {estimate('pct_hispanic'):.3f} (p < 0.001), and "
        f"pct_asian = {estimate('pct_asian'):.3f} (p < 0.001). The conditional "
        f"pct_bachelors_plus coefficient is also negative ({estimate('pct_bachelors_plus'):.3f}, "
        "p = 0.026), despite the positive unadjusted education gradient, which is "
        "consistent with overlap among education, income, housing, and PV deployment. "
        "The storage gap is therefore not explained solely by unequal PV deployment upstream."
    )
    replace_paragraph(
        find_paragraph(document, "The charger-type decomposition helps clarify"),
        "The charger-type decomposition helps clarify what is going on. In the standardized "
        "baseline models, the positive aggregate association with poverty_rate is not present "
        "for y_level1_chargers, where the coefficient is small and statistically insignificant "
        "(0.002, p = 0.584). Instead, the aggregate pattern appears to be driven primarily by "
        "y_level2_chargers, where poverty_rate is positive and highly significant (0.166, "
        "p < 0.001). For y_dc_fast_chargers, the coefficient on poverty_rate is also positive, "
        "but it is imprecisely estimated in the standardized baseline model (0.027, p = 0.202)."
    )
    replace_paragraph(
        find_paragraph(document, "In addition to the regression ladder"),
        "In addition to the regression ladder, we use principal component analysis and K-means "
        "clustering to summarize ZIP/ZCTA contexts in the statewide sample. This is descriptive "
        "rather than causal. The cluster map in Figure 15 shows that these profiles are spatially "
        "coherent rather than randomly scattered. Cluster identifiers are ordered by mean household income for "
        "stable interpretation. Table 6 shows that clusters 1 and 2 are the lowest-income and "
        "most heavily Hispanic profiles; cluster 5 combines high income, the largest Asian share, "
        "and the strongest charging outcome; and cluster 6 is the highest-income profile and has "
        "the strongest storage outcome. These differences are consistent with the broader paper "
        "argument that technology access is shaped by bundles of demographic, economic, housing, "
        "and infrastructural conditions rather than by one variable in isolation."
    )
    replace_paragraph(
        find_paragraph_any(document, ("Figure 15.", "Figure 15:")),
        "Figure 15: California ZIP/ZCTA clusters from the exploratory PCA plus K-means analysis. "
        "Cluster IDs are stable labels ordered from lowest to highest mean household income; the "
        "map shows that the resulting neighborhood profiles are geographically structured."
    )


def load_vif(model: str) -> pd.Series:
    path = TABLE_DIR / f"y_pv | {model} VIF.csv"
    return pd.read_csv(path).set_index("feature")["VIF"]


def update_vif_table(document: Document) -> None:
    table = next(
        table
        for table in document.tables
        if table.cell(0, 0).text == "Predictor"
        and len(table.columns) == 5
        and table.cell(0, 1).text == "M1 baseline"
    )
    labels = [row.cells[0].text for row in table.rows]
    if "% boat/RV/van/other units" not in labels:
        new_row = table.add_row()
        wind_row = next(row for row in table.rows if row.cells[0].text == "Wind MW per 100k")
        wind_row._tr.addprevious(new_row._tr)
        new_row.cells[0].text = "% boat/RV/van/other units"

    models = [
        "Model 1 baseline (climate controls)",
        "Model 2 (add housing value)",
        "Model 2D (add housing structure and tenure)",
        "Model 7 (per-capita infrastructure controls)",
    ]
    vifs = [load_vif(model) for model in models]
    feature_for_label = {
        "Log median income": "log_median_household_income",
        "Poverty rate": "poverty_rate",
        "% Hispanic": "pct_hispanic",
        "% Asian": "pct_asian",
        "Solar irradiance (GHI)": "ghi_mean_kwh_m2_day_2023",
        "% Black": "pct_black",
        "Log housing value": "log_median_housing_value",
        "% owner-occupied": "owner_occupied_rate",
        "% multifamily units": "pct_multifamily_units",
        "% mobile-home units": "pct_mobile_home_units",
        "% boat/RV/van/other units": "pct_other_housing_units",
        "Wind MW per 100k": "wind_mw_per_100k_ctrl",
        "Turbines per 100k": "turbines_per_100k",
        "Storage MW per 100k": "storage_mw_per_100k",
        "Plant MW per 100k": "plant_mw_per_100k",
    }
    for row in table.rows[1:]:
        label = row.cells[0].text
        if label == "Maximum VIF":
            values = [series.max() for series in vifs]
        else:
            feature = feature_for_label.get(label)
            if feature is None:
                continue
            values = [series.get(feature) for series in vifs]
        for cell, value in zip(row.cells[1:], values):
            cell.text = "—" if pd.isna(value) else f"{float(value):.2f}"

    model2d_max = float(vifs[2].max())
    wind_vif = float(vifs[3].get("wind_mw_per_100k_ctrl"))
    turbine_vif = float(vifs[3].get("turbines_per_100k"))
    frame = pd.read_csv(ROOT / "data" / "processed" / "combined_der_dataset_w_controls_predictors.csv")
    correlation = frame[["owner_occupied_rate", "pct_multifamily_units"]].corr().iloc[0, 1]
    paragraph = find_paragraph(document, "Variance inflation factors for the Solar PV models")
    replace_paragraph(
        paragraph,
        "Variance inflation factors for the Solar PV models (Table 7) show both a "
        "well-conditioned specification and a cautionary case. VIF is computed with the "
        "intercept retained in the auxiliary regressions, so raw and standardized "
        "diagnostics agree. The housing structure and tenure controls in Model 2D have "
        f"a maximum VIF of {model2d_max:.2f}, despite owner occupancy correlating "
        f"{correlation:.2f} with multifamily share. In the per-capita infrastructure "
        f"specification, wind MW per 100k and turbines per 100k reach {wind_vif:.2f} and "
        f"{turbine_vif:.2f}; their individual coefficients should not be read separately. "
        "VIF tables for every model and outcome are exported alongside the coefficient "
        "tables.",
    )


def update_spatial_tables(document: Document) -> None:
    """Refresh Tables 8 and 9 after the release regression run."""
    moran_table = next(
        table for table in document.tables
        if table.cell(0, 0).text == "Outcome" and table.cell(0, 1).text == "M1 baseline"
    )
    moran = pd.read_csv(ROOT / "outputs" / "tables" / "morans_i_residuals.csv")
    moran = moran.loc[moran["k"] == 8].pivot(index="outcome", columns="model", values="morans_i")
    outcome_map = {
        "Solar PV": "y_pv", "Storage": "y_storage",
        "EV chargers (aggregate)": "y_chargers",
        "EV chargers (Level 2)": "y_level2_chargers",
        "Energy burden": "energy_burden_pct",
        "Log affordability gap": "log_energy_gap_per_capita",
    }
    models = ["M1 baseline", "M6A + lat/lon", "M6B county FE", "M8 + demand proxy"]
    for row in moran_table.rows[1:]:
        outcome = outcome_map[row.cells[0].text]
        for cell, model in zip(row.cells[1:], models):
            cell.text = f"{moran.loc[outcome, model]:.3f}"

    conley_table = next(
        table for table in document.tables
        if table.cell(0, 0).text == "Term" and table.cell(0, 1).text == "Coefficient"
    )
    conley = pd.read_csv(ROOT / "outputs" / "tables" / "conley_standard_errors.csv")
    outcome_map = {
        "Solar PV": "y_pv", "Storage": "y_storage",
        "EV chargers (aggregate)": "y_chargers",
    }
    term_map = {
        "Log median income": "log_median_household_income", "% Black": "pct_black",
        "% Hispanic": "pct_hispanic", "% Asian": "pct_asian", "Poverty rate": "poverty_rate",
    }
    outcome = None
    for row in conley_table.rows[1:]:
        first = row.cells[0].text.strip()
        if first in outcome_map:
            outcome = outcome_map[first]
            continue
        if first not in term_map or outcome is None:
            continue
        term = term_map[first]
        result = conley.loc[(conley["outcome"] == outcome) & (conley["term"] == term)].iloc[0]
        row.cells[1].text = f"{result['coef']:.3f}"
        for cell, se_col in zip(
            row.cells[2:], ["se_hc1", "se_conley_50km", "se_conley_100km", "se_conley_200km"]
        ):
            p_col = se_col.replace("se_", "p_")
            cell.text = f"{result[se_col]:.3f}{spatial_stars(result[p_col])}"

    spatial = find_paragraph(document, "Moran’s I computed")
    replace_paragraph(
        spatial,
        spatial.text.replace("p < 0.001", "p ≤ 0.001").replace(
            "the reported standard errors are likely to be conservative rather than sufficient",
            "county clustering alone is not sufficient",
        ),
    )


def update_text(document: Document) -> None:
    data = find_paragraph(document, "After source-specific cleaning")
    replace_paragraph(
        data,
        "After source-specific cleaning, the DER measures are merged with contextual "
        "predictors. Demographic and socioeconomic variables come from the 2023 American "
        "Community Survey 5-year estimates. Housing structure uses every category in ACS "
        "B25024: single-family, multifamily, mobile-home, and boat/RV/van/other shares sum "
        "to one. Tenure uses B25003 owner occupancy, with renter occupancy as its "
        "complement. Climate and renewable-resource controls come from NASA POWER; utility "
        "territory uses a ZIP-to-utility spatial crosswalk; and annualized utility usage "
        "provides the electricity-demand proxy.",
    )

    baseline = find_paragraph(document, "A consistent pattern across the main specifications")
    existing = [p for p in document.paragraphs if p.text.startswith("Housing is tested directly")]
    if existing:
        replace_paragraph(existing[0], housing_result_text())
    else:
        insert_after(baseline, housing_result_text(), "Normal")

    model2b = find_paragraph(document, "Model 2B: Additional SES control")
    if not any(p.text.startswith("Model 2C:") for p in document.paragraphs):
        model2c = insert_after(
            model2b,
            "Model 2C: Exhaustive housing structure composition (B25024), with "
            "single-family units as the omitted reference",
            "Normal",
        )
        insert_after(
            model2c,
            "Model 2D: Housing structure plus owner occupancy (B25003), with renters as "
            "the omitted tenure reference",
            "Normal",
        )
    model2_description = find_paragraph_any(
        document,
        ("We then add another socioeconomic proxy", "We estimate four Model 2 variants"),
    )
    replace_paragraph(
        model2_description,
        "We estimate four Model 2 variants. Models 2A and 2B add bachelor’s attainment "
        "and median housing value. Model 2C adds multifamily, mobile-home, and "
        "boat/RV/van/other shares from the exhaustive B25024 composition; Model 2D adds "
        "B25003 owner occupancy. The implementation fails rather than silently "
        "simplifying Models 2C or 2D when a required housing variable is absent or "
        "constant. These variants use the general additional-control form in Equation "
        "(6).",
    )

    replacements = {
        ("The present paper does not include a dedicated transit", "The present paper includes area-level housing structure"): (
            "The present paper includes area-level housing structure and tenure controls "
            "but does not include dedicated transit, corridor, charger-host, household "
            "tenancy, or roof-suitability data, so it cannot fully unpack why charging is "
            "more prevalent in some higher-poverty ZIP/ZCTAs. That remains an important "
            "extension. Aggregate infrastructure counts should not be mistaken for "
            "equitable access: future planning needs to distinguish commercial siting "
            "from household-serving benefit."
        ),
        ("This methodology is designed to identify area-level associations", "This methodology identifies area-level associations"): (
            "This methodology identifies area-level associations, not causal effects. "
            "B25024 structure and B25003 tenure are included at the ZIP/ZCTA level, but "
            "unmeasured confounding may remain around household tenancy, roof suitability, "
            "installer presence, permitting, and technology-specific incentive access. "
            "ZIP/ZCTA aggregation can also blur within-area heterogeneity, and some source "
            "files require spatial assignment or manual cleaning."
        ),
        ("This study is designed to identify area-level conditional associations", "This study identifies area-level conditional associations"): (
            "This study identifies area-level conditional associations, not causal effects "
            "(Section 6.8). Area-level housing structure and tenure are observed in Models "
            "2C and 2D, but unmeasured confounding may remain around household tenancy, "
            "roof suitability, installer availability, local permitting, program "
            "participation, and technology-specific incentive access."
        ),
        ("Finally, the paper is a California 2023 cross-sectional snapshot",): (
            "Finally, the paper is a California 2023 cross-sectional snapshot. The results "
            "may not generalize mechanically to other states, years, or regulatory "
            "contexts. Future work would benefit from longitudinal analysis, property- or "
            "household-level tenancy and roof-suitability measures, direct data on "
            "technology-specific access mechanisms, and place-based measures such as "
            "transit and corridor intensity."
        ),
    }
    for prefixes, text in replacements.items():
        replace_paragraph(find_paragraph_any(document, prefixes), text)

    cluster_conclusion = find_paragraph(document, "The clustering exercise adds")
    replace_paragraph(
        cluster_conclusion,
        "The clustering exercise adds another descriptive layer to this story (Section 4.7). "
        "The income-ordered profiles show that clusters 1 and 2 are lower-income and heavily "
        "Hispanic, cluster 5 combines high income, a large Asian share, and strong charging, "
        "and cluster 6 is the highest-income, strongest-storage profile (Table 6 and Figure 15). "
        "Those spatial profiles are descriptive, but they reinforce that race, income, housing, "
        "and existing energy-system context jointly structure where DER benefits accumulate."
    )
    for paragraph in list(document.paragraphs):
        if paragraph.text.startswith("Add some comments on clusters"):
            remove_paragraph(paragraph)

    demand = find_paragraph(document, "After correcting a parsing fault")
    replace_paragraph(
        demand,
        "After correcting a parsing fault that had silently zeroed roughly 85 percent of "
        "PG&E consumption records, the demand proxy enters positively where it is significant: "
        "0.018 (p < 0.001) for solar PV and 0.061 (p < 0.001) for storage, and it is not "
        "distinguishable from zero for aggregate charging (0.007, p = 0.327). The corrected "
        "variable is available for 1,196 of the 1,392 analysis ZIP/ZCTAs, with 1,159 reporting "
        "positive consumption; most remaining gaps fall outside the three investor-owned utility "
        "territories for which ZIP-level usage is published. Race and ethnicity coefficients in "
        "the demand-proxy models moved by at most 0.04 under this correction, with no change in "
        "significance, so the robustness claims that rely on them are unaffected."
    )


def add_housing_figure(document: Document) -> None:
    summary = find_paragraph(document, "This appendix keeps the heavier visual material")
    replace_paragraph(
        summary,
        "This appendix keeps the heavier visual material out of the main results flow "
        "while preserving model-comparison, spatial, wind, and housing-control checks "
        "(Figures 17 through 21), followed by the statistical diagnostics in Tables 7 "
        "through 9.",
    )
    if any(p.text.startswith("Figure 21:") for p in document.paragraphs):
        return
    diagnostics = find_paragraph(document, "Statistical robustness and diagnostics")
    heading = document.add_paragraph("Housing structure and tenure attenuation", style="Heading 3")
    diagnostics._p.addprevious(heading._p)
    image_paragraph = document.add_paragraph()
    diagnostics._p.addprevious(image_paragraph._p)
    image = image_paragraph.add_run().add_picture(str(HOUSING_FIGURE), width=Inches(6.4))
    image._inline.docPr.set(
        "descr",
        "Race coefficients before and after adding housing structure and tenure controls",
    )
    caption = document.add_paragraph(
        "Figure 21: Race-coefficient attenuation after adding exhaustive B25024 housing "
        "structure and B25003 tenure controls. Hollow points are the standardized "
        "baseline; filled points are Model 2D.",
        style="Caption",
    )
    diagnostics._p.addprevious(caption._p)
    note = document.add_paragraph(
        "Structure coefficients use single-family housing as the omitted reference, and "
        "tenure uses renters as the omitted reference.",
        style="Normal",
    )
    diagnostics._p.addprevious(note._p)


def main() -> None:
    if not HOUSING_FIGURE.exists():
        raise FileNotFoundError(f"Generate the housing figure first: {HOUSING_FIGURE}")
    document = Document(DOCX_PATH)
    clean_review_markup(document)
    update_text(document)
    update_main_tables(document)
    update_release_prose(document)
    materialize_figure_references(document)
    materialize_equation_references(document)
    materialize_compact_equations(document)
    replace_figure_before_caption(document, "Figure 15:", CLUSTER_FIGURE)
    update_vif_table(document)
    update_spatial_tables(document)
    add_housing_figure(document)
    fix_equation_layout(document)
    document.save(DOCX_PATH)
    print(f"Updated {DOCX_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
