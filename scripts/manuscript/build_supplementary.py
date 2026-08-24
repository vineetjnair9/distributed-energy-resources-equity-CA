#!/usr/bin/env python3
"""Build the Supplementary Information document.

Written to be read rather than skimmed: thirteen numbered Notes, each carrying prose
with its figures and tables instead of a bare caption list. Most of the text is the
original appendix and Results material from the draft, restored (see ``si_content.py``);
the rest documents the analyses added for this revision.

Reuses the original paragraphs for anything carrying an image or a table, so the media
parts survive untouched, and generates the new tables from ``outputs/nature_revision/``.

Run from the repository root, after ``nature_revision_analysis.py``:

    python scripts/manuscript/build_supplementary.py

Item numbering is shared with the main text. ``build_manuscript.py`` writes the main
text's references against the same scheme; renumber in one place and you must renumber
in the other.
"""
from __future__ import annotations

import copy, os, re, shutil, sys, zipfile
import xml.etree.ElementTree as ET
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import si_content as SI
from spelling import americanize

W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
WP='{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
XS='{http://www.w3.org/XML/1998/namespace}space'
for _p,_u in [('w',W),('a',A),('r',R),('wp',WP)]: ET.register_namespace(_p,_u[1:-1])

ROOT=os.path.abspath(os.environ.get('REPO_ROOT',
      os.path.join(os.path.dirname(os.path.abspath(__file__)),'..','..')))
BUILD=os.path.join(ROOT,'build')
NR=os.path.join(ROOT,'outputs','nature_revision')
FIGS=os.path.join(ROOT,'outputs','standardized_figures')
SOURCE_DOCX=os.path.join(ROOT,'site','index.docx')
DEST_DOCX=os.path.join(ROOT,'site','supplementary_information.docx')

os.makedirs(BUILD,exist_ok=True)
SRC=os.path.join(BUILD,'_source')
if not os.path.exists(SRC):
    with zipfile.ZipFile(SOURCE_DOCX) as _z: _z.extractall(SRC)

_tree=ET.parse(os.path.join(SRC,'word','document.xml'))
_body=_tree.getroot().find(W+'body')
CH=list(_body)
NB=[]; add=NB.append

def para_text(p): return ''.join(x.text or '' for x in p.iter(W+'t'))

def declean(el):
    """Strip comment anchors and comment-reference runs from a copied element."""
    for parent in el.iter():
        for c in list(parent):
            if c.tag in (W+'commentRangeStart',W+'commentRangeEnd'): parent.remove(c)
            elif c.tag==W+'r' and c.find(W+'commentReference') is not None: parent.remove(c)
    return el

def keep(i): return declean(copy.deepcopy(CH[i]))

def P(style,text):
    p=ET.Element(W+'p')
    if style:
        pr=ET.SubElement(p,W+'pPr'); st=ET.SubElement(pr,W+'pStyle'); st.set(W+'val',style)
    r=ET.SubElement(p,W+'r'); tt=ET.SubElement(r,W+'t')
    tt.set(XS,'preserve'); tt.text=text
    return p

TBL_PROTO=CH[30]
def table(rows):
    tbl=ET.Element(W+'tbl')
    pr=TBL_PROTO.find(W+'tblPr')
    if pr is not None: tbl.append(copy.deepcopy(pr))
    g=ET.SubElement(tbl,W+'tblGrid')
    for _ in rows[0]: ET.SubElement(g,W+'gridCol')
    for ri,row in enumerate(rows):
        tr=ET.SubElement(tbl,W+'tr')
        for cell in row:
            tc=ET.SubElement(tr,W+'tc'); ET.SubElement(tc,W+'tcPr')
            p=ET.SubElement(tc,W+'p'); r=ET.SubElement(p,W+'r')
            if ri==0:
                rp=ET.SubElement(r,W+'rPr'); ET.SubElement(rp,W+'b')
            tt=ET.SubElement(r,W+'t'); tt.set(XS,'preserve'); tt.text=str(cell)
    return tbl

def fig(i,cap): add(keep(i)); add(P('Caption',cap))
def tab(cap,rows): add(P('Caption',cap)); add(table(rows))
def tab_src(cap,i): add(P('Caption',cap)); add(keep(i))
def note(i):
    title,paras=SI.NOTES[i]
    add(P('Heading1',title))
    for t in paras: add(P(None,t))

NICE={'y_pv':'Rooftop solar','y_storage':'Battery storage','y_chargers':'EV charging',
 'y_level2_chargers':'Level 2 chargers','y_dc_fast_chargers':'DC fast chargers',
 'y_level1_chargers':'Level 1 chargers','y_wind_mw':'Wind capacity',
 'energy_burden_pct':'Energy burden','log_energy_gap_per_capita':'Affordability gap',
 'log_median_household_income':'Log income','pct_black':'Black share',
 'pct_hispanic':'Hispanic share','pct_asian':'Asian share','poverty_rate':'Poverty rate'}
def csv(n): return pd.read_csv(os.path.join(NR,n))
def stars(p): return '***' if p<0.001 else ('**' if p<0.01 else ('*' if p<0.05 else ''))

# ---------------------------------------------------------------- front matter
add(P('Title','Supplementary Information'))
add(P(None,'Inequity in distributed energy adoption is technology-specific, not uniform'))
add(P(None,SI.INTRO))

# ---- S1 model series ----
note(0)
fig(44,'Supplementary Fig. S1 | Rooftop solar: every coefficient across the model series. One '
       'panel per coefficient, with 95 per cent confidence bands; filled markers denote p < 0.05.')
fig(54,'Supplementary Fig. S2 | Battery storage: every coefficient across the model series. The '
       'race and ethnicity bands stay clear of zero across every specification, which is the '
       'basis for treating storage as the strongest result in the study.')
fig(60,'Supplementary Fig. S3 | Aggregate EV charging: every coefficient across the model '
       'series. Poverty rate is the only consistently signed predictor; the race terms cross zero.')
tab('Supplementary Table S1 | Specifications in the model series. Each varies one control block '
    'against the fixed core of income, race and ethnicity shares, poverty rate and the '
    'outcome-specific resource control.',
 [['Label','Deviation from the core'],
  ['M1','Core specification'],
  ['M2A / M2B','Adds educational attainment, or median housing value'],
  ['M2C / M2D','Adds exhaustive housing structure, or structure and tenure'],
  ['M3A / M3B / M3C','Resource control substituted: degree days, mean temperature, or global horizontal irradiance'],
  ['M4 / M4R','Adds centred income-by-race interactions; M4R drops the poverty control'],
  ['M5 / M5C','Adds utility fixed effects; M5C re-estimates the core with county-clustered standard errors'],
  ['M6A / M6B','Replaces the resource control with latitude and longitude, or with county fixed effects'],
  ['M7 / M7PC','Adds installed infrastructure capacity, in levels or per capita'],
  ['M8','Adds the logged electricity-demand proxy'],
  ['M9','Storage only: adds local solar deployment'],
  ['M9A','Energy-burden outcomes: adds DER deployment terms']])

# ---- S2 nested cumulative ladder ----
note(1)
fig(27,'Supplementary Fig. S4 | Coefficients across the nested cumulative ladder, C1 to C5. '
       'Standardised predictors; frozen common sample of 1,160 ZCTAs at every rung, so movement '
       'reflects added controls rather than sample composition.')
L=csv('ladder_c1_c5.csv'); L=L[L.outcome.isin(['y_pv','y_storage','y_chargers'])]
rows=[['Outcome','Term','C1','C2','C3','C4','C5']]
for o in ['y_pv','y_storage','y_chargers']:
    for tm in ['log_median_household_income','pct_black','pct_hispanic','pct_asian']:
        g=L[(L.outcome==o)&(L.term==tm)].set_index('rung')
        cells=['%.3f%s'%(g.loc[k,'beta'],stars(g.loc[k,'p'])) if k in g.index else '—'
               for k in ['C1','C2','C3','C4','C5']]
        rows.append([NICE[o],NICE[tm]]+cells)
tab('Supplementary Table S5 | Nested cumulative ladder, C1 to C5. Standardised predictors; '
    'frozen common sample of 1,160 ZCTAs and 42 county clusters, identical at every rung. C5 uses '
    'county-clustered standard errors. *p < 0.05, **p < 0.01, ***p < 0.001.',rows)
tab_src('Supplementary Table S6 | Battery storage conditional on local solar deployment. This '
        'model conditions on a parallel outcome of the same predictors and is not a '
        'better-identified estimate of the storage gap.',51)

# ---- S3 specification curve ----
note(2)
fig(37,'Supplementary Fig. S5 | Specification curve for rooftop solar, battery storage and '
       'aggregate EV charging. Each panel sorts the focal coefficient across all 48 valid '
       'combinations of the six confounder blocks, with the block-membership matrix below.')

# ---- S4 Oster bounds ----
note(3)
O=pd.read_csv(os.path.join(ROOT,'outputs','tables','oster_delta.csv'))
tab('Supplementary Table S7 | Oster δ bounds, core specification versus the saturated '
    'confounder model. δ ≥ 1 is the conventional robustness threshold. Negative or very '
    'large values indicate that adding controls increased rather than attenuated the '
    'coefficient, so the bound does not apply.',
 [['Outcome','Term','β restricted','β full','R² restricted','R² full','δ']]+
 [[NICE.get(r.outcome,r.outcome),NICE.get(r.term,r.term),'%.3f'%r.beta_restricted,
   '%.3f'%r.beta_full,'%.3f'%r.r2_restricted,'%.3f'%r.r2_full,'%.2f'%r.delta]
  for r in O[O.outcome.isin(['y_pv','y_storage','y_chargers'])].itertuples()])

# ---- S5 functional form, weighting, spatial ----
note(4)
Pm=csv('ppml_negbin.csv'); Pm=Pm[Pm.model=='PPML']
tab('Supplementary Table S8 | Poisson pseudo-maximum-likelihood models on raw counts with a '
    'population offset. Fitted directly to counts rather than to a log(1 + rate) transformation, '
    'so structural zeros are handled without transformation. HC1 standard errors.',
 [['Outcome','Term','Coefficient','Std. error','p','N']]+
 [[r.outcome,NICE.get(r.term,r.term),'%.3f'%r.coef,'%.3f'%r.se,'%.3f'%r.p,str(int(r.N))]
  for r in Pm.itertuples()])
Wt=csv('population_weighted.csv')
tab('Supplementary Table S9 | Population-weighted versus unweighted baseline estimates. '
    'Standardised predictors; weights are ZCTA total population.',
 [['Outcome','Term','Unweighted','p','Weighted','p','N']]+
 [[NICE[r.outcome],NICE[r.term],'%.3f'%r.beta_unweighted,'%.3f'%r.p_unw,
   '%.3f'%r.beta_weighted,'%.3f'%r.p_w,str(int(r.N))] for r in Wt.itertuples()])
S=csv('spatial_error_model.csv')
tab('Supplementary Table S10 | Spatial error model, maximum likelihood, eight-nearest-neighbour '
    'weights. Standardised predictors; λ is the spatial autoregressive error parameter.',
 [['Outcome','Term','Coefficient','Std. error','p','λ','N']]+
 [[NICE[r.outcome],NICE[r.term],'%.3f'%r.beta_SEM,'%.3f'%r.se_SEM,'%.3f'%r.p_SEM,
   '%.2f'%r.lambda_,str(int(r.N))] for r in S.itertuples()])

# ---- S6 housing ----
note(5)
tab('Supplementary Table S2 | Race-coefficient attenuation after adding housing structure and '
    'tenure. Standardised coefficients; baseline versus the specification adding exhaustive ACS '
    'B25024 structure shares and B25003 owner occupancy.',
 [['Outcome','Term','Baseline','+ structure & tenure','Attenuation'],
  ['Rooftop solar','Black share','−0.049','−0.023','53%'],
  ['Rooftop solar','Asian share','−0.082','−0.051','38%'],
  ['Battery storage','Black share','−0.106','−0.059','44%'],
  ['Battery storage','Asian share','−0.252','−0.171','32%'],
  ['Aggregate charging','Asian share','0.058','−0.008','not distinguishable from zero']])

# ---- S7 descriptive gradients ----
note(6)
fig(81,'Supplementary Fig. S6 | Descriptive LOWESS gradients between combined non-white share '
       'and the three headline outcomes. Shaded ribbons are 95 per cent bootstrap confidence '
       'intervals; these panels do not condition on the full control set.')
fig(83,'Supplementary Fig. S7 | Descriptive LOWESS gradients between bachelor’s-degree '
       'attainment and the three headline outcomes.')
fig(136,'Supplementary Fig. S8 | Distributions of raw deployment rates together with median '
        'household income and poverty rate.')

# ---- S8 typology ----
note(7)
fig(88,'Supplementary Fig. S9 | California ZCTA clusters from exploratory principal-component '
       'and K-means analysis. Cluster identifiers run from lowest to highest mean household income.')
tab_src('Supplementary Table S16 | Cluster-profile means from the principal-component and '
        'K-means assignment.',91)

# ---- S9 burden and affordability ----
note(8)
fig(72,'Supplementary Fig. S10 | Descriptive LOWESS gradients between neighbourhood '
       'characteristics and the two energy-affordability outcomes.')
fig(74,'Supplementary Fig. S11 | DER predictors in the energy-burden and affordability-gap '
       'models. Points are coefficient estimates; horizontal lines are 95 per cent confidence '
       'intervals.')
fig(76,'Supplementary Fig. S12 | Coefficient stability for the energy-burden and '
       'affordability-gap outcomes across the model series.')
Af=csv('affordability_outliers.csv')
tab('Supplementary Table S11 | Energy-affordability models with and without the 134 ZCTAs whose '
    'per-capita affordability gap exceeds $1,000. Standardised predictors.',
 [['Outcome','Sample','Term','Coefficient','p','N','R²']]+
 [[NICE.get(r.outcome,r.outcome),r.sample,NICE.get(r.term,r.term),'%.3f'%r.beta,'%.3f'%r.p,
   str(int(r.N)),'%.3f'%r.R2] for r in Af.itertuples()])

# ---- S10 wind ----
note(9)
fig(95,'Supplementary Fig. S13 | Wind capacity: coefficients across the model series, included '
       'as contextual material. Wind capacity is zero in 97.5 per cent of ZCTAs.')

# ---- S11 predictive comparison and spatial output ----
note(10)
fig(220,'Supplementary Fig. S14 | Battery storage: cross-validated predictive comparison of '
        'linear and nonlinear models. Error bars are ±1 s.d. across five cross-validation folds.')
fig(224,'Supplementary Fig. S15 | Aggregate EV charging: cross-validated predictive comparison of '
        'linear and nonlinear models.')
fig(232,'Supplementary Fig. S16 | Wind capacity: cross-validated predictive comparison of linear '
        'and nonlinear models.')
fig(228,'Supplementary Fig. S17 | Rooftop solar: observed adoption and baseline standardised '
        'residuals by ZCTA. The clustering visible in the residual panel is what the Moran’s '
        'I statistics in Note S12 quantify.')

# ---- S12 diagnostics ----
note(11)
tab_src('Supplementary Table S17 | Variance inflation factors across four rooftop solar '
        'specifications. Values above about 10 indicate a coefficient is not separately '
        'identified.',244)
tab_src('Supplementary Table S18 | Moran’s I on OLS residuals, eight-nearest-neighbour '
        'weights. Every entry is significant at p < 0.001 by permutation test.',248)
tab_src('Supplementary Table S19 | Conley spatial HAC standard errors, Bartlett kernel over '
        'great-circle distance. Coefficients are identical across columns; only the standard '
        'errors differ.',252)

# ---- S13 data construction ----
note(12)
Z=csv('zero_shares.csv')
tab('Supplementary Table S13 | Share of ZCTAs reporting zero deployment, by outcome. Level 1 '
    'charging and wind capacity are too sparse to support a coefficient at this geography.',
 [['Outcome','Source variable','N','% zero','Median','Maximum']]+
 [[NICE.get(r.outcome,r.outcome),r.raw_var,str(int(r.N)),'%.1f'%r.pct_zero,'%.2f'%r.median,
   '%.1f'%r.max] for r in Z.itertuples()])
St=csv('storage_by_sector.csv')
piv={(r.sample,r.term):r for r in St.itertuples()}
srows=[['Term','All sectors','p','Residential only','p']]
for tm in ['log_median_household_income','pct_black','pct_hispanic','pct_asian','poverty_rate']:
    a=piv.get(('all sectors',tm)); b=piv.get(('residential only',tm))
    srows.append([NICE.get(tm,tm),
                  '%.3f'%a.coef if a else '—','%.3f'%a.p if a else '—',
                  '%.3f'%b.coef if b else '—','%.3f'%b.p if b else '—'])
tab('Supplementary Table S12 | Storage sensitivity by customer sector. The California Energy '
    'Commission records are 193,070 residential, 3,211 commercial and 291 utility. Standardised '
    'predictors; models re-fitted on capacity aggregated from all sectors and from residential '
    'records only.',srows)
tab_src('Supplementary Table S14 | Variable definitions for the main outcomes, predictors and '
        'controls.',129)
tab_src('Supplementary Table S15 | Summary statistics for the main outcomes and key predictors '
        'in the analysis sample.',139)
sd=csv('sd_conversion.csv')
tab('Supplementary Table S3 | Predictor means and standard deviations on the analysis sample, '
    'for converting standardised coefficients to percentage-point changes.',
 [['Predictor','Mean','Standard deviation','One s.d. in percentage points']]+
 [[NICE.get(r.predictor,r.predictor),'%.3f'%r.mean,'%.3f'%r.sd,
   ('—' if str(r.predictor).startswith('log') else '%.1f'%r.one_sd_in_pp)]
  for r in sd.itertuples()])
tab_src('Supplementary Table S4 | EV charger subtype comparison in baseline models.',65)

add(copy.deepcopy(CH[269]))

# ---------------------------------------------------------------- package
def _set(p,s):
    ts=[x for x in p.iter(W+'t')]
    if not ts: return
    ts[0].text=s; ts[0].set(XS,'preserve')
    for x in ts[1:]: x.text=''

# British to American spelling, applied to the assembled text so words split across
# adjacent string literals in si_content.py are caught too.
for _el in NB:
    for _p in _el.iter(W+'p'):
        _s=para_text(_p); _n=americanize(_s)
        if _n!=_s: _set(_p,_n)

for c in list(_body): _body.remove(c)
for el in NB: _body.append(el)
print('SI body elements: %d'%len(list(_body)))

OUT=os.path.join(BUILD,'si')
if os.path.exists(OUT): shutil.rmtree(OUT)
shutil.copytree(SRC,OUT)
for f in ('comments.xml','commentsExtended.xml','commentsExtensible.xml','commentsIds.xml','people.xml'):
    fp=os.path.join(OUT,'word',f)
    if os.path.exists(fp): os.remove(fp)
_tree.write(os.path.join(OUT,'word','document.xml'),xml_declaration=True,encoding='UTF-8')

def spec_curve_panel():
    """Stack the three per-outcome specification-curve exports into one SI panel."""
    from PIL import Image
    ims=[Image.open(os.path.join(FIGS,'spec_curve_y_%s.png'%n)).convert('RGB')
         for n in ('pv','storage','chargers')]
    w=max(i.width for i in ims)
    ims=[i.resize((w,int(i.height*w/i.width)),Image.LANCZOS) if i.width!=w else i for i in ims]
    gap=40
    out=Image.new('RGB',(w,sum(i.height for i in ims)+gap*(len(ims)-1)),'white')
    y=0
    for i in ims:
        out.paste(i,(0,y)); y+=i.height+gap
    path=os.path.join(BUILD,'spec_curve_panel.png'); out.save(path,dpi=(300,300))
    return path

def swap_image(rel_id,media_name,source_png,width_emu=5943600):
    """Replace the bytes behind one drawing and reset its extent for the new aspect ratio."""
    from PIL import Image
    shutil.copy(source_png,os.path.join(OUT,'word','media',media_name))
    iw,ih=Image.open(os.path.join(OUT,'word','media',media_name)).size
    cy=int(width_emu*ih/iw)
    f=os.path.join(OUT,'word','document.xml'); tree=ET.parse(f)
    for dr in tree.getroot().iter(W+'drawing'):
        blip=dr.find('.//'+A+'blip')
        if blip is None or blip.get(R+'embed')!=rel_id: continue
        ext=dr.find('.//'+WP+'extent')
        if ext is not None: ext.set('cx',str(width_emu)); ext.set('cy',str(cy))
        for e in dr.iter(A+'ext'): e.set('cx',str(width_emu)); e.set('cy',str(cy))
    tree.write(f,xml_declaration=True,encoding='UTF-8')

# Fig S4 reuses the drawing that carried the old income-gradient figure and Fig S5 the
# cross-outcome dot-whisker; both are repointed at the new exports.
swap_image('rId10','image1.png',os.path.join(FIGS,'coefficient_path_c1_c5.png'))
swap_image('rId12','image3.png',spec_curve_panel())

rp=os.path.join(OUT,'word','_rels','document.xml.rels'); s=open(rp,encoding='utf8').read()
s=re.sub(r'<Relationship[^>]*Target="comments[^"]*"[^>]*/>','',s)
s=re.sub(r'<Relationship[^>]*Target="people\.xml"[^>]*/>','',s)
open(rp,'w',encoding='utf8').write(s)
cp=os.path.join(OUT,'[Content_Types].xml'); s=open(cp,encoding='utf8').read()
for part in ('comments','commentsExtended','commentsExtensible','commentsIds','people'):
    s=re.sub(r'<Override[^>]*PartName="/word/%s\.xml"[^>]*/>'%part,'',s)
open(cp,'w',encoding='utf8').write(s)

if os.path.exists(DEST_DOCX): os.remove(DEST_DOCX)
zf=zipfile.ZipFile(DEST_DOCX,'w',zipfile.ZIP_DEFLATED)
zf.write(cp,'[Content_Types].xml')
for r,_d,files in os.walk(OUT):
    for f in files:
        full=os.path.join(r,f); rel=os.path.relpath(full,OUT)
        if rel!='[Content_Types].xml': zf.write(full,rel)
zf.close()

import docx
doc=docx.Document(DEST_DOCX)
txt='\n'.join(p.text for p in doc.paragraphs)
words=sum(len(p.text.split()) for p in doc.paragraphs if p.style.name!='Caption')
print('wrote %s (%.1f MB)'%(DEST_DOCX,os.path.getsize(DEST_DOCX)/1e6))
print('  %d Notes, %d figures, %d tables, %d words of prose'%(
  len(re.findall(r'Supplementary Note S\d+ \|',txt)),
  len(re.findall(r'Supplementary Fig\. S\d+ \|',txt)),
  len(doc.tables),words))
